#!/usr/bin/env python3
"""
Propuestas de renovación BSE a partir del export "Exportar resultados a Excel"
de Consulta de Cotizaciones del Portal del Asesor. No toca BSE.

Uso:
  python3 bse_propuesta.py ConsultaCotizaciones.xls                 # resumen + Excel de trabajo (todas)
  python3 bse_propuesta.py ConsultaCotizaciones.xls --pct 3         # aumento objetivo sobre Premio Póliza (default 3)
  python3 bse_propuesta.py ConsultaCotizaciones.xls --cliente vilomon --pdf   # PDF para un cliente

Lógica (validada contra BSE el 19/10/2026 con 5 pólizas: desvío < 0,03 %):
  C = Premio Cotización, P = Premio Póliza
  objetivo = P * (1 + pct/100)
  bonificación x = (1 - objetivo / C) * 100, redondeada a 2 decimales
  premio propuesto = C * (1 - x/100)
  si x < 0 (C ya está por debajo del objetivo): no se bonifica, queda C.

Ojo: una cotización ya editada aparece en el export con Premio Cotización VACÍO.
Exportar ANTES de editar y guardar ese archivo como base.
"""
import argparse, sys, os, datetime
import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def load(path):
    b = xlrd.open_workbook(path)
    s = b.sheet_by_index(0)
    H = [str(h).strip() for h in s.row_values(0)]
    rows = []
    for r in range(1, s.nrows):
        d = dict(zip(H, s.row_values(r)))
        num = lambda v: float(v) if v not in ("", None) else None
        txt = lambda v: str(int(v)) if isinstance(v, float) else str(v).strip()
        rows.append({
            "cotizacion": txt(d["Cotización"]),
            "poliza": txt(d["Póliza"]) if d["Póliza"] != "" else "",
            "cliente": str(d["Nombre del Cliente"]).strip(),
            "desde": d["Fecha Desde"], "hasta": d["Fecha Hasta"],
            "moneda": d["Moneda"], "plan": d["Plan de pago"],
            "C": num(d["Premio Cotización"]), "P": num(d["Premio Póliza"]),
            "error": d["Error"], "pauta": d["En Pauta"],
        })
    return rows


def calc(r, pct):
    C, P = r["C"], r["P"]
    if C is None or P is None or C <= 0:
        r.update(estado="sin datos", x=None, objetivo=None, propuesto=None, var=None)
        return r
    obj = P * (1 + pct / 100)
    x = round((1 - obj / C) * 100, 2)
    if x < 0:
        r.update(estado="no alcanza (queda cotización BSE)", x=0.0, objetivo=obj, propuesto=C)
    elif x > 50:
        r.update(estado="bonificación > 50 % (inválida en BSE)", x=x, objetivo=obj, propuesto=None)
    else:
        r.update(estado="ajustar" + (" (fuera de pauta >10 %)" if x > 10 else ""), x=x, objetivo=obj,
                 propuesto=round(C * (1 - x / 100), 2))
    r["var"] = (r["propuesto"] / P - 1) * 100 if r.get("propuesto") else None
    return r


def money(v):
    return "" if v is None else "$ " + f"{v:,.0f}".replace(",", ".")


def write_xlsx(rows, out, pct):
    wb = Workbook(); ws = wb.active; ws.title = "Propuestas"
    hdr = ["Cliente", "Cotización", "Póliza", "Desde", "Hasta", "Plan de pago", "Premio póliza (P)",
           "Premio cotización BSE (C)", f"Objetivo P+{pct:g}%", "Bonificación a cargar %", "Premio propuesto",
           "Variación vs póliza %", "Estado", "Error BSE", "En pauta BSE"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in sorted(rows, key=lambda r: (r["cliente"], r["cotizacion"])):
        ws.append([r["cliente"], int(r["cotizacion"]), int(r["poliza"]) if r["poliza"] else None, r["desde"], r["hasta"],
                   r["plan"], r["P"], r["C"], r["objetivo"], r["x"], r["propuesto"],
                   round(r["var"], 2) if r["var"] is not None else None, r["estado"], r["error"], r["pauta"]])
    for col, w in zip(range(1, len(hdr) + 1), [38, 12, 11, 11, 11, 22, 14, 16, 14, 13, 14, 12, 34, 8, 8]):
        ws.column_dimensions[get_column_letter(col)].width = w
    for row in ws.iter_rows(min_row=2):
        for i in (6, 7, 8, 10):
            row[i].number_format = '#,##0'
        row[9].number_format = '0.00'; row[11].number_format = '0.00'
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    wb.save(out)


def write_pdf(rows, out, pct):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    ss = getSampleStyleSheet(); navy = colors.HexColor("#1F3A5F")
    h = ParagraphStyle("h", parent=ss["Title"], fontSize=18, textColor=navy, alignment=0, spaceAfter=4)
    sub = ParagraphStyle("s", parent=ss["Normal"], fontSize=10, textColor=colors.HexColor("#444444"), leading=14)
    cell = ParagraphStyle("c", parent=ss["Normal"], fontSize=9, leading=11)
    cellr = ParagraphStyle("cr", parent=cell, alignment=2)
    note = ParagraphStyle("n", parent=ss["Normal"], fontSize=8.5, textColor=colors.HexColor("#555555"), leading=12)
    cliente = rows[0]["cliente"]
    desde = sorted({r["desde"] for r in rows}); hasta = sorted({r["hasta"] for r in rows})
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm,
                            bottomMargin=15*mm, title=f"Propuesta de renovación - {cliente}")
    story = [Paragraph("Propuesta de renovación de seguros", h),
             Paragraph(f"<b>Cliente:</b> {cliente} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Aseguradora:</b> Banco de Seguros del Estado", sub),
             Paragraph(f"<b>Vigencia:</b> {desde[0] if len(desde)==1 else ', '.join(desde)} al {hasta[0] if len(hasta)==1 else ', '.join(hasta)}"
                       f" &nbsp;&nbsp;|&nbsp;&nbsp; <b>Fecha:</b> {datetime.date.today():%d/%m/%Y}", sub), Spacer(1, 8)]
    hdr = ["Póliza", "Plan de pago", "Premio actual", "Premio renovación", "Variación"]
    data = [[Paragraph(f"<b>{x}</b>", ParagraphStyle("hh", parent=cell, textColor=colors.white)) for x in hdr]]
    ta = tn = 0
    for r in sorted(rows, key=lambda r: r["poliza"]):
        if r["propuesto"] is None:
            continue
        ta += r["P"]; tn += r["propuesto"]
        data.append([r["poliza"], r["plan"], money(r["P"]), Paragraph(f"<b>{money(r['propuesto'])}</b>", cellr),
                     f"{r['var']:+.1f} %".replace(".", ",")])
    data.append(["", Paragraph("<b>Total</b>", cell), money(ta), Paragraph(f"<b>{money(tn)}</b>", cellr),
                 f"{(tn/ta-1)*100:+.1f} %".replace(".", ",") if ta else ""])
    t = Table(data, colWidths=[30*mm, 60*mm, 40*mm, 45*mm, 30*mm], repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), navy), ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F2F5F9")]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E3E9F1")), ("LINEABOVE", (0, -1), (-1, -1), 0.8, navy),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D2DE")),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story += [t, Spacer(1, 10), Paragraph(
        "Importes en pesos uruguayos, impuestos incluidos. Propuesta estimada sobre la cotización de renovación vigente "
        "a la fecha; el importe final puede diferir en algunos pesos por redondeo y queda sujeto a la emisión de la póliza "
        "y a las condiciones generales y particulares de cada una.", note)]
    doc.build(story)
    return ta, tn


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xls"); ap.add_argument("--pct", type=float, default=3.0)
    ap.add_argument("--cliente"); ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--out", default=".")
    a = ap.parse_args()
    rows = [calc(r, a.pct) for r in load(a.xls)]
    if a.cliente:
        rows = [r for r in rows if a.cliente.lower() in r["cliente"].lower()]
        if not rows:
            sys.exit("No hay cotizaciones para ese cliente")
    import collections
    est = collections.Counter(r["estado"] for r in rows)
    print(f"Cotizaciones: {len(rows)}  Clientes: {len({r['cliente'] for r in rows})}")
    for k, v in est.most_common():
        print(f"  {v:4d}  {k}")
    stamp = datetime.date.today().strftime("%Y-%m-%d")
    xlsx = os.path.join(a.out, f"propuestas_{(a.cliente or 'todas').replace(' ', '_')}_{stamp}.xlsx")
    write_xlsx(rows, xlsx, a.pct); print("Excel:", xlsx)
    if a.pdf:
        if len({r["cliente"] for r in rows}) != 1:
            sys.exit("Para PDF filtrá un solo cliente con --cliente")
        pdf = os.path.join(a.out, f"Propuesta_{rows[0]['cliente'].split()[0]}_{stamp}.pdf")
        ta, tn = write_pdf(rows, pdf, a.pct); print("PDF:", pdf, money(ta), "->", money(tn))


if __name__ == "__main__":
    main()
